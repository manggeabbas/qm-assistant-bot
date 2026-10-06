"""PPTX extractor (text + tables per slide)."""

from __future__ import annotations

from pathlib import Path

from qm_training.material.base import MaterialBlock, MaterialDocument


def read_pptx(path: Path) -> MaterialDocument:
    from pptx import Presentation

    presentation = Presentation(str(path))
    blocks: list[MaterialBlock] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame:
            title_text = (title_shape.text_frame.text or "").strip()
            if title_text:
                parts.append(title_text)  # title first: main title candidate
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if shape.has_text_frame:
                text = (shape.text_frame.text or "").strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        blocks.append(MaterialBlock("slide", "\n".join(parts), meta={"slide": index}))
    return MaterialDocument(source=str(path), format="pptx", blocks=blocks)

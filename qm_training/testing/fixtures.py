"""Build material fixtures for tests (test utilities)."""

from __future__ import annotations

from pathlib import Path

from docx import Document


def build_docx(path: Path) -> Path:
    document = Document()
    document.add_heading("Prosedur Keselamatan Laboratorium", level=1)
    document.add_paragraph("Gunakan alat pelindung diri sebelum masuk area laboratorium.")
    document.add_paragraph("Selalu periksa peralatan sebelum digunakan.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Bahaya"
    table.cell(0, 1).text = "Tindakan"
    table.cell(1, 0).text = "Tumpahan kimia"
    table.cell(1, 1).text = "Gunakan spill kit"
    document.save(str(path))
    return path


def build_mandarin_docx(path: Path) -> Path:
    document = Document()
    document.add_heading("质量管理部培训签到表", level=1)
    document.add_paragraph("培训内容：安全操作规程")
    document.save(str(path))
    return path


def build_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Emu

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Briefing K3 Laboratorium"
    box = slide.shapes.add_textbox(Emu(500000), Emu(2000000), Emu(8000000), Emu(1500000))
    box.text_frame.text = "Selalu periksa peralatan sebelum digunakan."
    presentation.save(str(path))
    return path


def build_bilingual_pptx(path: Path) -> Path:
    """Slide 1 with an Indonesian main title and a Mandarin subtitle."""
    from pptx import Presentation
    from pptx.util import Emu

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = '"2024.08.28" Peninjauan Kecelakaan Tersayat Sampel'
    box = slide.shapes.add_textbox(Emu(500000), Emu(3000000), Emu(8000000), Emu(1500000))
    box.text_frame.text = "2024.08.28 样品割伤事故回顾"
    presentation.save(str(path))
    return path


def build_xlsx(path: Path) -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Checklist"
    sheet.append(["No", "Item"])
    sheet.append([1, "Helm"])
    sheet.append([2, "Sarung tangan"])
    schedule = workbook.create_sheet("Jadwal")
    schedule.append(["Tanggal", "Kegiatan"])
    schedule.append(["03-09-2026", "Briefing"])
    workbook.save(str(path))
    return path


def build_material_fixtures(out_dir: Path, convert_to_legacy: bool = True) -> dict[str, Path]:
    """Return a mapping extension -> fixture path for all 7 formats."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    fixtures: dict[str, Path] = {}
    fixtures[".docx"] = build_docx(out_dir / "material.docx")
    fixtures[".pptx"] = build_pptx(out_dir / "material.pptx")
    fixtures[".xlsx"] = build_xlsx(out_dir / "material.xlsx")

    if convert_to_legacy:
        from qm_training.core.libreoffice import convert

        fixtures[".pdf"] = convert(fixtures[".docx"], "pdf", out_dir)
        fixtures[".doc"] = convert(fixtures[".docx"], "doc", out_dir)
        fixtures[".ppt"] = convert(fixtures[".pptx"], "ppt", out_dir)
        fixtures[".xls"] = convert(fixtures[".xlsx"], "xls", out_dir)
    return fixtures

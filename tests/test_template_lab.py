"""Automated structural/pagination tests for the Template Laboratory (Fase 1).

Jalankan:
    .venv/bin/python -m unittest discover -s tests -v
"""

from __future__ import annotations

import hashlib
import json
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from template_lab import dummy_data
from template_lab.convert_pdf import convert_to_pdf, find_soffice
from template_lab.generator import cells, generate, load_map, paragraphs, rows
from template_lab.paths import MASTER_TEMPLATE, OUTPUT_DIR
from template_lab.render_pages import pdf_page_count, pdf_pages_text

TEST_DIR = OUTPUT_DIR / "tests"
SCENARIOS: dict[str, Path] = {}
PDF_PATH: Path | None = None


def text_of(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def tables_of(doc: Document) -> list:
    return doc.element.body.findall(qn("w:tbl"))


def tables_with(doc: Document, marker: str) -> list:
    return [table for table in tables_of(doc) if marker in text_of(table)]


def body_children(doc: Document) -> list:
    return list(doc.element.body)


def setUpModule() -> None:
    TEST_DIR.mkdir(parents=True, exist_ok=True)
    SCENARIOS["full"] = generate(dummy_data.dummy_data(15, 3), out_dir=TEST_DIR, output_name="test_full")
    SCENARIOS["skip"] = generate(dummy_data.dummy_data(None, 1), out_dir=TEST_DIR, output_name="test_skip")
    SCENARIOS["zero"] = generate(dummy_data.dummy_data(15, 0), out_dir=TEST_DIR, output_name="test_zero")
    SCENARIOS["many"] = generate(dummy_data.dummy_data(15, 5), out_dir=TEST_DIR, output_name="test_many")

    global PDF_PATH
    if find_soffice() is not None:
        try:
            PDF_PATH = convert_to_pdf(SCENARIOS["full"], TEST_DIR)
        except Exception:  # pragma: no cover - PDF conversion is environment dependent
            PDF_PATH = None


class TestMasterAndMap(unittest.TestCase):
    def test_master_is_read_only(self) -> None:
        digest = hashlib.sha256(MASTER_TEMPLATE.read_bytes()).hexdigest()
        template = load_map()
        self.assertEqual(digest, template["generated_from"]["master_sha256"])

    def test_map_targets_resolve_on_master(self) -> None:
        template = load_map()
        doc = Document(str(MASTER_TEMPLATE))
        tbls = tables_of(doc)

        def check(table_index: int, row: int, cell: int, paragraph: int, run: int | None = None) -> None:
            tr = rows(tbls[table_index])[row]
            tc = cells(tr)[cell]
            ps = paragraphs(tc)
            self.assertLess(paragraph, len(ps), f"paragraph {paragraph} hilang")
            if run is not None:
                self.assertLess(run, len(ps[paragraph].findall(qn("w:r"))), f"run {run} hilang")

        for field in template["dynamic_fields"]:
            targets = field.get("targets")
            if not isinstance(targets, list):
                continue
            for target in targets:
                check(target["table_index"], target["row"], target["cell"], target["paragraph"])
                source = target.get("style_source")
                if source:
                    check(target["table_index"], source["row"], source["cell"], source["paragraph"], source["run"])

        placeholder = template["t03a"]["question_placeholder"]
        for question in template["t03a"]["questions"]["rows"]:
            check(2, question["row"], 0, question["placeholder_paragraph"])
            cell = rows(tbls[2])[question["row"]].findall(qn("w:tc"))[0]
            paragraph = cell.findall(qn("w:p"))[question["placeholder_paragraph"]]
            text = "".join(node.text or "" for node in paragraph.iter(qn("w:t")))
            self.assertIn(placeholder, text)
        for answer in template["t03a"]["answer_areas"]["rows"]:
            check(2, answer["row"], 0, answer["answer_area_paragraph"])


class TestT01A(unittest.TestCase):
    def setUp(self) -> None:
        self.doc = Document(str(SCENARIOS["full"]))
        self.table = tables_with(self.doc, "Tema Pelatihan")[0]

    def test_personnel_rows_are_15_and_numbered_1_to_15(self) -> None:
        data_rows = rows(self.table)[5:20]
        self.assertEqual(len(data_rows), 15)
        numbers = [text_of(paragraphs(cells(row)[0])[0]).strip() for row in data_rows]
        self.assertEqual(numbers, [str(i) for i in range(1, 16)])

    def test_dynamic_fields_filled(self) -> None:
        text = text_of(self.table)
        for expected in (
            "Keselamatan Kerja di Area Laboratorium Quality Management",
            "2026-QM-LZ-09-03",
            "03 September 2026",
            "Ruang Training Quality Management",
            "Budi Santoso",
            "15人",
        ):
            self.assertIn(expected, text)

    def test_participant_cells_are_blank(self) -> None:
        for row in rows(self.table)[5:20]:
            for cell_index in (1, 2, 3, 5, 6, 7):
                self.assertEqual(text_of(cells(row)[cell_index]).strip(), "")


class TestPersonnelSkip(unittest.TestCase):
    def test_skip_leaves_cell_empty(self) -> None:
        doc = Document(str(SCENARIOS["skip"]))
        table = tables_with(doc, "Personil Absen")[0]
        value_cell = cells(rows(table)[3])[3]
        text = text_of(value_cell).strip()
        self.assertEqual(text, "")
        self.assertNotIn("0人", text)
        self.assertNotIn("N/A", text)


class TestT03A(unittest.TestCase):
    def setUp(self) -> None:
        self.doc = Document(str(SCENARIOS["full"]))
        self.table = tables_with(self.doc, "PENCATATAN PERTANYAAN")[0]

    def test_identity_values(self) -> None:
        def cell_text(row: int, cell: int) -> str:
            return text_of(cells(rows(self.table)[row])[cell])

        self.assertEqual(cell_text(0, 1).strip(), "")       # Nama kosong
        self.assertEqual(cell_text(0, 3).strip(), "")       # NIK kosong
        self.assertEqual(cell_text(0, 5).strip(), "QM")
        self.assertEqual(cell_text(1, 1).strip(), "Budi Santoso")
        self.assertEqual(cell_text(1, 3).strip(), "03 September 2026")
        self.assertEqual(cell_text(1, 5).strip(), "QMYWI")
        self.assertIn("Keselamatan Kerja", cell_text(2, 1))
        self.assertEqual(cell_text(2, 3).strip(), "")        # NILAI kosong

    def test_exactly_five_questions_present(self) -> None:
        text = text_of(self.table)
        for index in range(1, 6):
            self.assertIn(f"PERTANYAAN {index}", text)
        self.assertNotIn("PERTANYAAN 6", text)
        self.assertIn("Apa tujuan utama briefing keselamatan kerja", text)

    def test_answer_areas_are_blank(self) -> None:
        for row_index in (5, 7, 9, 11, 13):
            answer_area = paragraphs(cells(rows(self.table)[row_index])[0])[1]
            self.assertEqual(text_of(answer_area).strip(), "")


class TestAnswerKey(unittest.TestCase):
    def setUp(self) -> None:
        self.doc = Document(str(SCENARIOS["full"]))

    def test_answer_key_present_after_t03a_with_page_break(self) -> None:
        children = body_children(self.doc)
        t03a_index = next(
            i for i, el in enumerate(children)
            if el.tag == qn("w:tbl") and "PENCATATAN PERTANYAAN" in text_of(el)
        )
        heading_index = next(
            i for i, el in enumerate(children)
            if el.tag == qn("w:p") and "KUNCI JAWABAN" in text_of(el)
        )
        self.assertGreater(heading_index, t03a_index)

        between = children[t03a_index + 1:heading_index + 1]

        def starts_new_page(element) -> bool:
            for br in element.iter(qn("w:br")):
                if br.get(qn("w:type")) == "page":
                    return True
            for pbb in element.iter(qn("w:pageBreakBefore")):
                if (pbb.get(qn("w:val")) or "1") not in ("0", "false"):
                    return True
            return False

        page_break_found = any(starts_new_page(el) for el in between)
        self.assertTrue(page_break_found, "halaman kunci jawaban harus dimulai di halaman baru")

        text = text_of(children[heading_index])
        self.assertIn("KUNCI JAWABAN", text)
        self.assertIn("答案", text)

    def test_answer_key_contains_five_answers(self) -> None:
        full_text = "\n".join(text_of(el) for el in body_children(self.doc))
        start = full_text.index("KUNCI JAWABAN")
        answer_key_text = full_text[start:]
        self.assertGreaterEqual(answer_key_text.count("JAWABAN 答"), 5)


class TestStaticAndPhotos(unittest.TestCase):
    def test_form_codes_preserved(self) -> None:
        text = "\n".join(text_of(el) for el in body_children(Document(str(SCENARIOS["full"]))))
        for code in ("QMZ15002-T01A", "QMZ15002-T02A", "QMZ15002-T03A"):
            self.assertIn(code, text)

    def test_no_master_sample_photos_in_output(self) -> None:
        with zipfile.ZipFile(SCENARIOS["full"]) as archive:
            names = set(archive.namelist())
        self.assertFalse(any(name in names for name in (
            "word/media/image3.jpeg", "word/media/image4.jpeg", "word/media/image5.jpeg"
        )), "foto contoh master tidak boleh ikut ke output")

    def test_photo_pages_and_frame(self) -> None:
        template = load_map()
        frame = template["t02a"]["cell_frame"]
        max_w = frame["inner_width_twips"] * 635
        max_h = (frame["row_height_twips"] - 700) * 635

        many = Document(str(SCENARIOS["many"]))
        t02a_tables = tables_with(many, "Foto pertama")
        self.assertEqual(len(t02a_tables), 3)  # 5 foto -> 3 halaman T02A

        inline = [d for d in many.element.body.iter(qn("w:drawing")) if d.find(qn("wp:inline")) is not None]
        photo_drawings = [d for d in inline if d.find(".//" + qn("a:blip")) is not None]
        self.assertGreaterEqual(len(photo_drawings), 5)
        for drawing in photo_drawings:
            extent = drawing.find(qn("wp:inline")).find(qn("wp:extent"))
            self.assertLessEqual(int(extent.get("cx")), max_w)
            self.assertLessEqual(int(extent.get("cy")), max_h)

    def test_zero_photos_still_valid(self) -> None:
        doc = Document(str(SCENARIOS["zero"]))
        t02a_tables = tables_with(doc, "Foto pertama")
        self.assertEqual(len(t02a_tables), 1)
        self.assertEqual(len(rows(t02a_tables[0])), 2)

    def test_drawing_ids_are_unique(self) -> None:
        from collections import Counter

        for scenario in SCENARIOS.values():
            doc = Document(str(scenario))
            doc_pr = [node.get("id") for node in doc.element.body.iter(qn("wp:docPr"))]
            duplicates = [value for value, count in Counter(doc_pr).items() if count > 1]
            self.assertEqual(duplicates, [], f"docPr id ganda pada {scenario.name}")

    def test_photo_count_matches_uploaded(self) -> None:
        with zipfile.ZipFile(SCENARIOS["many"]) as archive:
            names = [n for n in archive.namelist() if n.startswith("word/media/") and n not in (
                "word/media/image1.png", "word/media/image2.png"
            )]
        self.assertEqual(len(names), 5)


class TestPdf(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if PDF_PATH is None:
            raise unittest.SkipTest("LibreOffice/PDF tidak tersedia")

    def test_page_count_and_content(self) -> None:
        # T01A + T03A + 2 halaman T02A (3 foto) + kunci jawaban = 5
        self.assertEqual(pdf_page_count(PDF_PATH), 5)

    def test_answer_key_on_last_page(self) -> None:
        pages = pdf_pages_text(PDF_PATH)
        self.assertIn("KUNCI JAWABAN", pages[-1])

    def test_t03a_not_split_across_pages(self) -> None:
        pages = pdf_pages_text(PDF_PATH)
        t03a_pages = [text for text in pages if "PERTANYAAN 1" in text]
        self.assertEqual(len(t03a_pages), 1)
        self.assertIn("PERTANYAAN 5", t03a_pages[0])
        self.assertGreaterEqual(t03a_pages[0].count("JAWABAN"), 5)


if __name__ == "__main__":
    unittest.main()

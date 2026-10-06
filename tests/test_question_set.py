"""QuestionSet lifecycle: generated once, previewed once, frozen on confirm.

Checklist:
1  one material -> one QuestionSet
2  one QuestionSet -> one preview
3  no second preview
4  state WAITING_QUESTION_CONFIRMATION after preview
5  confirm does not create a new QuestionSet
6  duplicate confirmation ignored
7  QuestionSet stored in session
8  DOCX uses the session QuestionSet
9  DOCX contains all five questions
10 placeholder removed
11 no new T03A row
12 no new T03A table
13 no border added by the generator
14 participant answer areas stay blank
"""

from __future__ import annotations

import shutil
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from qm_training.ai.mock import MockAIProvider
from qm_training.ai.service import AIService
from qm_training.bot import SessionStore, Workflow
from qm_training.bot.adapters import MockTelegramAdapter
from qm_training.bot.states import State
from qm_training.bot.workflow import BTN_CONFIRM, BTN_USE_THEME
from qm_training.core.config import Settings
from qm_training.document.engine import generate as real_generate
from qm_training.document.pdf import convert_to_pdf, find_soffice
from qm_training.paths import MASTER_TEMPLATE
from qm_training.testing.fixtures import build_docx
from qm_training.validation.pipeline import validate_document
from qm_training.validation.structure import compare_t03a_structure

BASE = Path("output/phase10_questions")
MATERIAL: Path


def setUpModule() -> None:
    global MATERIAL
    if find_soffice() is None:
        raise unittest.SkipTest("LibreOffice tidak tersedia untuk e2e")
    if BASE.exists():
        shutil.rmtree(BASE)
    BASE.mkdir(parents=True, exist_ok=True)
    MATERIAL = build_docx(BASE / "material.docx")


def text_of(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def t03a(doc: Document):
    for table in doc.element.body.findall(qn("w:tbl")):
        if "PENCATATAN PERTANYAAN" in text_of(table):
            return table
    raise AssertionError("T03A tidak ditemukan")


class TestQuestionSetLifecycle(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        case = BASE / "case"
        store = SessionStore(case / "sessions")
        workflow = Workflow(
            store=store,
            ai_service=AIService(MockAIProvider()),
            settings=Settings(ai_provider="mock"),
            generate_docx=real_generate,
            convert_pdf=convert_to_pdf,
            validate=validate_document,
        )
        cls.adapter = MockTelegramAdapter(workflow)
        cls.store = store

        cls.adapter.receive("1", "/new_training", update_id=1)
        cls.adapter.receive("1", document_path=MATERIAL, update_id=2)
        cls.adapter.receive("1", BTN_USE_THEME, update_id=3)
        cls.adapter.receive("1", "2026-QM-9-17-WA", update_id=4)
        cls.adapter.receive("1", "2026/09/15", update_id=5)
        cls.adapter.receive("1", "Area Mesin", update_id=6)
        cls.adapter.receive("1", "Ayub", update_id=7)
        cls.adapter.receive("1", "REGU A", update_id=8)
        cls.adapter.receive("1", "15", update_id=9)

        cls.preview_messages = cls.adapter.receive("1", BTN_CONFIRM, update_id=10)
        cls.session_after_preview = store.get("1")
        cls.preview_questions = list(cls.session_after_preview.questions)

        # stray text while awaiting confirmation must not resend a preview
        cls.stray_messages = cls.adapter.receive("1", "apa saja", update_id=11)

        # confirm the QuestionSet
        cls.confirm_messages = cls.adapter.receive("1", BTN_CONFIRM, update_id=12)
        # duplicate confirmation (same update id) and a second one (new id)
        cls.duplicate_same = cls.adapter.receive("1", BTN_CONFIRM, update_id=12)
        cls.duplicate_new = cls.adapter.receive("1", BTN_CONFIRM, update_id=13)

        cls.finish_messages = cls.adapter.receive("1", "✅ Selesai", update_id=14)
        cls.session = store.get("1")

    # 1 -------------------------------------------------------------------
    def test_1_one_material_one_question_set(self) -> None:
        self.assertEqual(self.session.generate_questions_calls, 1)
        self.assertTrue(self.session.question_set_created)
        self.assertEqual(len(self.session.questions), 5)

    # 2, 3 ----------------------------------------------------------------
    def test_2_3_one_preview_only(self) -> None:
        previews = [m for m in self.preview_messages if "Preview Pertanyaan" in m.text]
        self.assertEqual(len(previews), 1)
        self.assertEqual(self.session.preview_sent, 1)
        stray_previews = [m for m in self.stray_messages if "Preview Pertanyaan" in m.text]
        self.assertEqual(stray_previews, [])
        self.assertEqual(self.session.preview_sent, 1)

    # 4 -------------------------------------------------------------------
    def test_4_state_after_preview(self) -> None:
        self.assertEqual(self.session_after_preview.state, State.WAITING_QUESTION_CONFIRMATION.value)

    # 5, 6 ----------------------------------------------------------------
    def test_5_6_confirm_does_not_regenerate(self) -> None:
        self.assertEqual(self.session.generate_questions_calls, 1)
        self.assertEqual(self.session.questions, self.preview_questions)
        self.assertTrue(self.session.question_confirmed)
        self.assertEqual(self.duplicate_same, [])
        self.assertEqual(self.duplicate_new, [])
        self.assertEqual(self.session.generate_questions_calls, 1)

    # 7 -------------------------------------------------------------------
    def test_7_question_set_in_session(self) -> None:
        self.assertEqual(len(self.session.questions), 5)
        for item in self.session.questions:
            self.assertTrue(item["question"].strip())
            self.assertTrue(item["answer"].strip())

    # 8, 9, 10 ------------------------------------------------------------
    def test_8_9_10_docx_uses_session_questions(self) -> None:
        docx = Path(self.session.output_docx)
        self.assertTrue(docx.exists())
        doc = Document(str(docx))
        full = text_of(doc.element.body)
        for item in self.session.questions:
            self.assertIn(item["question"], full, f"pertanyaan hilang di DOCX: {item['question'][:40]}")
        self.assertNotIn("TEMPATKAN PERTANYAAN DISINI", full)
        # preview == session == docx
        for index, item in enumerate(self.preview_questions, start=1):
            self.assertIn(item["question"], full, f"preview Q{index} != DOCX")

    # 11, 12, 13 ----------------------------------------------------------
    def test_11_12_13_t03a_structure_preserved(self) -> None:
        doc = Document(str(Path(self.session.output_docx)))
        master = Document(str(MASTER_TEMPLATE))
        self.assertEqual(len(t03a(doc).findall(qn("w:tr"))), len(t03a(master).findall(qn("w:tr"))))
        self.assertEqual(
            len([t for t in doc.element.body.findall(qn("w:tbl")) if "PENCATATAN PERTANYAAN" in text_of(t)]),
            1,
        )
        report = compare_t03a_structure(Path(self.session.output_docx), MASTER_TEMPLATE)
        self.assertEqual(report.errors, [], report.summary())

    # 14 ------------------------------------------------------------------
    def test_14_answer_areas_blank(self) -> None:
        doc = Document(str(Path(self.session.output_docx)))
        rows = t03a(doc).findall(qn("w:tr"))
        for index in range(1, 6):
            answer_paragraphs = rows[(index - 1) * 2 + 5].findall(qn("w:tc"))[0].findall(qn("w:p"))
            self.assertEqual(text_of(answer_paragraphs[1]).strip(), "")

    def test_completed_with_documents(self) -> None:
        self.assertEqual(self.session.state, State.COMPLETED.value)
        self.assertEqual(len(self.finish_messages[-1].documents), 2)

    def test_regenerate_is_explicit_and_new_set(self) -> None:
        # Regeneration is only allowed as an explicit user action.
        case = BASE / "regen"
        store = SessionStore(case / "sessions")
        wf = Workflow(
            store=store,
            ai_service=AIService(MockAIProvider()),
            settings=Settings(ai_provider="mock"),
            convert_pdf=lambda *a, **k: Path(self.session.output_pdf),
            validate=lambda d, p=None: validate_document(d, None),
        )
        adapter = MockTelegramAdapter(wf)
        adapter.receive("1", "/new_training", update_id=1)
        adapter.receive("1", document_path=MATERIAL, update_id=2)
        adapter.receive("1", BTN_USE_THEME, update_id=3)
        for i, value in enumerate(["A", "D", "L", "T", "S", "15"], start=4):
            adapter.receive("1", value, update_id=i)
        adapter.receive("1", BTN_CONFIRM, update_id=10)
        before = store.get("1").generate_questions_calls
        adapter.receive("1", "🔄 Generate Ulang", update_id=11)
        after = store.get("1").generate_questions_calls
        self.assertEqual(before, 1)
        self.assertEqual(after, 2)


if __name__ == "__main__":
    unittest.main()

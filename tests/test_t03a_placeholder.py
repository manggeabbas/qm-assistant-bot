"""T03A placeholder replacement: text-only, structure preserved.

Covers the new master template contract:
- questions replace the `TEMPATKAN PERTANYAAN DISINI` placeholder;
- no new row / table / cell / border / merge is created;
- T03A structure stays identical to the master.
"""

from __future__ import annotations

import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from qm_training.document import generate
from qm_training.document.pdf import convert_to_pdf, find_soffice
from qm_training.document.render import pdf_page_count, render_pdf_to_png
from qm_training.document.schema import QAPair
from qm_training.paths import MASTER_TEMPLATE
from qm_training.testing.datasets import make_photos
from qm_training.validation import validate_document
from qm_training.validation.structure import compare_t03a_structure
from qm_training.document.template import load_map

OUT = Path("output/phase8_t03a")

QUESTIONS = [
    QAPair("Pada tanggal berapa saja kecelakaan jari terjepit terjadi?", "3 Januari 2026."),
    QAPair("Di lokasi atau ruang apa kecelakaan kedua terjadi?", "Area mesin penggiling."),
    QAPair("Apa penyebab kecelakaan dari faktor manusia?", "Kurang konsentrasi."),
    QAPair("Apa penyebab kecelakaan dari faktor fisik?", "Pelindung mesin tidak terpasang."),
    QAPair("Apa langkah perbaikan terkait batang penyangga mesin penggiling?", "Periksa dan pasang kembali pengunci."),
]


def text_of(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def t03a(doc: Document):
    for table in doc.element.body.findall(qn("w:tbl")):
        if "PENCATATAN PERTANYAAN" in text_of(table):
            return table
    raise AssertionError("T03A tidak ditemukan")


def load_dataset():
    from qm_training.document.schema import TrainingData

    return TrainingData(
        theme="Pencegahan Kecelakaan Jari Terjepit",
        archive_number="2026-QM-9-17-WA",
        training_date="2026/09/15",
        location="Area Mesin Penggiling",
        trainer="Ayub",
        shift_group="REGU A",
        personnel_count=15,
        questions=list(QUESTIONS),
        photos=make_photos(3, OUT / "photos", prefix="t03a"),
    )


class TestT03APlaceholder(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        OUT.mkdir(parents=True, exist_ok=True)
        cls.docx = generate(load_dataset(), out_dir=OUT, output_name="t03a_case")
        cls.master = Document(str(MASTER_TEMPLATE))
        cls.generated = Document(str(cls.docx))

    def test_1_placeholder_removed(self) -> None:
        self.assertNotIn("TEMPATKAN PERTANYAAN DISINI", text_of(self.generated.element.body))

    def test_2_to_6_each_question_replaces_its_placeholder(self) -> None:
        rows = t03a(self.generated).findall(qn("w:tr"))
        for index, qa in enumerate(QUESTIONS, start=1):
            paragraph = rows[(index - 1) * 2 + 4].findall(qn("w:tc"))[0].findall(qn("w:p"))[0]
            text = text_of(paragraph)
            self.assertIn(f"PERTANYAAN {index}", text)
            self.assertIn(qa.question, text)
            self.assertNotIn("TEMPATKAN PERTANYAAN DISINI", text)

    def test_chinese_labels_preserved(self) -> None:
        rows = t03a(self.generated).findall(qn("w:tr"))
        for index in range(1, 6):
            paragraphs = rows[(index - 1) * 2 + 4].findall(qn("w:tc"))[0].findall(qn("w:p"))
            self.assertIn(f"问 题 {index}", text_of(paragraphs[1]))

    def test_7_8_no_extra_row_or_table_for_questions(self) -> None:
        self.assertEqual(len(t03a(self.generated).findall(qn("w:tr"))), 14)
        gen_tables = len(self.generated.element.body.findall(qn("w:tbl")))
        # master has 3 tables; only T02A may be duplicated for photos.
        self.assertGreaterEqual(gen_tables, 3)

    def test_9_to_12_structure_matches_master(self) -> None:
        report = compare_t03a_structure(self.docx, MASTER_TEMPLATE)
        self.assertEqual(report.errors, [], report.summary())

    def test_13_borders_not_created_by_generator(self) -> None:
        report = compare_t03a_structure(self.docx, MASTER_TEMPLATE)
        self.assertEqual([i for i in report.issues if "BORDER" in i.code], [])

    def test_13b_question_area_keeps_only_outer_horizontal_lines(self) -> None:
        rows_el = t03a(self.generated).findall(qn("w:tr"))

        def borders_of(row):
            table_ex = row.find(qn("w:tblPrEx"))
            return table_ex.find(qn("w:tblBorders")) if table_ex is not None else None

        # outer top edge of the question block (below the banner) is kept
        top_edge = borders_of(rows_el[3])
        self.assertEqual(top_edge.find(qn("w:bottom")).get(qn("w:val")), "single")

        # inner horizontal lines are hidden
        for row in rows_el[4:-1]:
            borders = borders_of(row)
            for tag in ("w:top", "w:bottom"):
                self.assertEqual(borders.find(qn(tag)).get(qn("w:val")), "none")

        # outer bottom edge of the table is kept
        bottom_edge = borders_of(rows_el[-1])
        self.assertEqual(bottom_edge.find(qn("w:top")).get(qn("w:val")), "none")
        self.assertEqual(bottom_edge.find(qn("w:bottom")).get(qn("w:val")), "single")

        # vertical border preserved
        table_borders = t03a(self.generated).find(qn("w:tblPr")).find(qn("w:tblBorders"))
        self.assertEqual(table_borders.find(qn("w:insideV")).get(qn("w:val")), "single")

    def test_14_answer_label_present(self) -> None:
        # 5 on T03A (participant answer areas) + 5 on the KUNCI JAWABAN page.
        self.assertEqual(text_of(t03a(self.generated)).count("JAWABAN 答"), 5)
        self.assertEqual(text_of(self.generated.element.body).count("JAWABAN 答"), 10)

    def test_15_answer_areas_blank(self) -> None:
        rows = t03a(self.generated).findall(qn("w:tr"))
        for index in range(1, 6):
            answer_paragraphs = rows[(index - 1) * 2 + 5].findall(qn("w:tc"))[0].findall(qn("w:p"))
            self.assertEqual(text_of(answer_paragraphs[1]).strip(), "")

    def test_16_question_labels_preserved(self) -> None:
        full = text_of(self.generated.element.body)
        for index in range(1, 6):
            self.assertIn(f"PERTANYAAN {index}", full)

    def test_17_logo_geometry_originates_from_template(self) -> None:
        def anchor_signatures(doc):
            signatures = []
            for anchor in doc.element.body.iter(qn("wp:anchor")):
                extent = anchor.find(qn("wp:extent"))
                def offset(axis):
                    position = anchor.find(qn(f"wp:{axis}"))
                    node = position.find(qn("wp:posOffset")) if position is not None else None
                    return node.text if node is not None else None

                signatures.append(
                    (
                        extent.get("cx") if extent is not None else None,
                        extent.get("cy") if extent is not None else None,
                        offset("positionH"),
                        offset("positionV"),
                    )
                )
            return signatures

        master_sigs = anchor_signatures(self.master)
        generated_sigs = anchor_signatures(self.generated)
        # Generator must never invent new logo geometry.
        self.assertTrue(set(generated_sigs).issubset(set(master_sigs)))
        # T01A logos (first table) keep exactly the master geometry.
        def table_anchors(doc):
            table = doc.element.body.findall(qn("w:tbl"))[0]
            return [
                (
                    a.find(qn("wp:extent")).get("cx"),
                    a.find(qn("wp:extent")).get("cy"),
                )
                for a in table.iter(qn("wp:anchor"))
            ]

        self.assertEqual(sorted(map(tuple, table_anchors(self.generated))), sorted(map(tuple, table_anchors(self.master))))

    def test_validation_pipeline_passes(self) -> None:
        report = validate_document(self.docx, None)
        self.assertEqual(report.errors, [], report.summary())

    def test_map_uses_placeholder_mode(self) -> None:
        template = load_map()
        self.assertEqual(template["t03a"]["question_placeholder"], "TEMPATKAN PERTANYAAN DISINI")
        for row in template["t03a"]["questions"]["rows"]:
            self.assertEqual(row["mode"], "replace_placeholder")

    def test_map_body_blocks_point_to_expected_elements(self) -> None:
        template = load_map()
        children = list(Document(str(MASTER_TEMPLATE)).element.body)
        t03a = children[template["t03a"]["body_block"]]
        self.assertEqual(t03a.tag, qn("w:tbl"))
        self.assertIn("PENCATATAN PERTANYAAN", text_of(t03a))
        t02a = children[template["t02a"]["body_block"]]
        self.assertEqual(t02a.tag, qn("w:tbl"))
        self.assertIn("Foto pertama", text_of(t02a))
        logo = children[template["answer_key"]["logo"]["source_body_block"]]
        logo_text = text_of(logo)
        self.assertTrue(logo_text.strip() == "" and len(list(logo.iter(qn("wp:anchor")))) == 2)


class TestT03AVisual(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if find_soffice() is None:
            raise unittest.SkipTest("LibreOffice tidak tersedia")
        cls.pdf = convert_to_pdf(OUT / "t03a_case.docx", OUT)
        cls.pngs = render_pdf_to_png(cls.pdf, OUT / "png")

    def test_pdf_and_render(self) -> None:
        self.assertGreaterEqual(pdf_page_count(self.pdf), 4)
        self.assertTrue(any(p.name.endswith(".png") for p in self.pngs))

    def test_t03a_page_has_questions(self) -> None:
        from qm_training.document.render import pdf_pages_text

        pages = pdf_pages_text(self.pdf)
        t03a_pages = [p for p in pages if "PERTANYAAN 1" in p]
        self.assertEqual(len(t03a_pages), 1)
        page = t03a_pages[0]
        for qa in QUESTIONS:
            self.assertIn(qa.question[:25], page)
        self.assertNotIn("TEMPATKAN PERTANYAAN DISINI", page)


if __name__ == "__main__":
    unittest.main()

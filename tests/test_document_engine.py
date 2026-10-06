"""Phase 2 tests: reusable document engine, schema and datasets."""

from __future__ import annotations

import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from qm_training.core.errors import SchemaError
from qm_training.document import generate
from qm_training.document.render import pdf_page_count
from qm_training.document.pdf import convert_to_pdf, find_soffice
from qm_training.document.schema import QAPair, TrainingData
from qm_training.testing.datasets import dataset

OUT = Path("output/phase2")


def text_of(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def tables_of(doc: Document) -> list:
    return doc.element.body.findall(qn("w:tbl"))


def tables_with(doc: Document, marker: str) -> list:
    return [t for t in tables_of(doc) if marker in text_of(t)]


def rows_of(table) -> list:
    return table.findall(qn("w:tr"))


class TestSchema(unittest.TestCase):
    def _base(self) -> TrainingData:
        return dataset(personnel_count=15, photo_count=0, photo_dir=OUT / "photos")

    def test_valid_dataset_passes(self) -> None:
        data = self._base()
        self.assertIs(data.validate(), data)

    def test_missing_field_rejected(self) -> None:
        data = self._base()
        data.location = "  "
        with self.assertRaises(SchemaError):
            data.validate()

    def test_wrong_question_count_rejected(self) -> None:
        data = self._base()
        data.questions = data.questions[:4]
        with self.assertRaises(SchemaError):
            data.validate()

    def test_empty_answer_rejected(self) -> None:
        data = self._base()
        data.questions[0] = QAPair("pertanyaan", "   ")
        with self.assertRaises(SchemaError):
            data.validate()

    def test_negative_personnel_rejected(self) -> None:
        data = self._base()
        data.personnel_count = -1
        with self.assertRaises(SchemaError):
            data.validate()

    def test_roundtrip_dict(self) -> None:
        original = self._base()
        restored = TrainingData.from_dict(original.to_dict())
        self.assertEqual(restored.to_dict(), original.to_dict())

    def test_filename_format(self) -> None:
        data = self._base()
        self.assertEqual(data.output_basename(), "2026-QM-LZ-09-03_REGU-A")


class TestEngineDatasets(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        OUT.mkdir(parents=True, exist_ok=True)
        photo_dir = OUT / "photos"
        cls.docs: dict[str, Path] = {}
        specs = {
            "personnel_null": dict(personnel_count=None, photo_count=0),
            "personnel_15": dict(personnel_count=15, photo_count=0),
            "photos_0": dict(personnel_count=15, photo_count=0),
            "photos_1": dict(personnel_count=15, photo_count=1),
            "photos_2": dict(personnel_count=15, photo_count=2),
            "photos_3": dict(personnel_count=15, photo_count=3),
            "photos_5": dict(personnel_count=15, photo_count=5),
            "short_q": dict(personnel_count=15, photo_count=0),
            "long_q": dict(personnel_count=15, photo_count=0, use_long_questions=True),
        }
        for name, kwargs in specs.items():
            data = dataset(photo_dir=photo_dir, **kwargs)
            cls.docs[name] = generate(data, out_dir=OUT, output_name=name)

    def test_all_documents_generated(self) -> None:
        self.assertEqual(len(self.docs), 9)
        for path in self.docs.values():
            self.assertTrue(path.exists())
            self.assertGreater(path.stat().st_size, 5000)

    def test_t01a_always_15_rows(self) -> None:
        for name in ("personnel_null", "personnel_15", "photos_5", "long_q"):
            doc = Document(str(self.docs[name]))
            table = tables_with(doc, "Tema Pelatihan")[0]
            numbers = [text_of(rows_of(table)[i].findall(qn("w:tc"))[0]).strip() for i in range(5, 20)]
            self.assertEqual(numbers, [str(i) for i in range(1, 16)], name)

    def test_personnel_null_vs_15(self) -> None:
        null_doc = Document(str(self.docs["personnel_null"]))
        fifteen_doc = Document(str(self.docs["personnel_15"]))
        null_cell = text_of(rows_of(tables_with(null_doc, "Personil Absen")[0])[3].findall(qn("w:tc"))[3]).strip()
        fifteen_cell = text_of(rows_of(tables_with(fifteen_doc, "Personil Absen")[0])[3].findall(qn("w:tc"))[3]).strip()
        self.assertEqual(null_cell, "")
        self.assertEqual(fifteen_cell, "15人")

    def test_photo_pages_scale(self) -> None:
        expected = {"photos_0": 1, "photos_1": 1, "photos_2": 1, "photos_3": 2, "photos_5": 3}
        for name, count in expected.items():
            doc = Document(str(self.docs[name]))
            self.assertEqual(len(tables_with(doc, "Foto pertama")), count, name)

    def test_questions_and_blank_answers(self) -> None:
        doc = Document(str(self.docs["long_q"]))
        table = tables_with(doc, "PENCATATAN PERTANYAAN")[0]
        text = text_of(table)
        for i in range(1, 6):
            self.assertIn(f"PERTANYAAN {i}", text)
        for row_index in (5, 7, 9, 11, 13):
            cell = rows_of(table)[row_index].findall(qn("w:tc"))[0]
            answer_area = cell.findall(qn("w:p"))[1]
            self.assertEqual(text_of(answer_area).strip(), "")

    def test_answer_key_present(self) -> None:
        doc = Document(str(self.docs["photos_3"]))
        full = "\n".join(text_of(el) for el in doc.element.body)
        self.assertIn("KUNCI JAWABAN", full)
        self.assertNotIn("QMZ15002-T04A", full)

    def test_no_sample_photos(self) -> None:
        with zipfile.ZipFile(self.docs["photos_5"]) as archive:
            names = set(archive.namelist())
        for banned in ("word/media/image3.jpeg", "word/media/image4.jpeg", "word/media/image5.jpeg"):
            self.assertNotIn(banned, names)

    def test_t03a_topic_tight_left(self) -> None:
        # Topik T03A (bagian yang diblok user): harus rata kiri + rapat,
        # tidak renggang.
        doc = Document(str(self.docs["personnel_15"]))
        table = tables_with(doc, "PENCATATAN PERTANYAAN")[0]
        cell = rows_of(table)[2].findall(qn("w:tc"))[1]
        paragraph = cell.findall(qn("w:p"))[0]
        self.assertIn("Keselamatan Kerja", text_of(paragraph))
        p_pr = paragraph.find(qn("w:pPr"))
        jc = p_pr.find(qn("w:jc"))
        self.assertEqual(jc.get(qn("w:val")), "left")
        spacing = p_pr.find(qn("w:spacing"))
        self.assertEqual(spacing.get(qn("w:before")), "0")
        self.assertEqual(spacing.get(qn("w:after")), "0")

    def test_photo_frames(self) -> None:
        doc = Document(str(self.docs["photos_5"]))
        inline = [d for d in doc.element.body.iter(qn("w:drawing")) if d.find(qn("wp:inline")) is not None]
        self.assertGreaterEqual(len(inline), 5)
        max_w = 8712 * 635
        max_h = (5472 - 700) * 635
        for drawing in inline:
            extent = drawing.find(qn("wp:inline")).find(qn("wp:extent"))
            self.assertLessEqual(int(extent.get("cx")), max_w)
            self.assertLessEqual(int(extent.get("cy")), max_h)

    def test_drawing_ids_unique(self) -> None:
        from collections import Counter

        for path in self.docs.values():
            doc = Document(str(path))
            ids = [n.get("id") for n in doc.element.body.iter(qn("wp:docPr"))]
            self.assertEqual([k for k, v in Counter(ids).items() if v > 1], [])


class TestEnginePdf(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if find_soffice() is None:
            raise unittest.SkipTest("LibreOffice tidak tersedia")
        cls.pdf_0 = convert_to_pdf(OUT / "photos_0.docx", OUT)
        cls.pdf_3 = convert_to_pdf(OUT / "photos_3.docx", OUT)
        cls.pdf_5 = convert_to_pdf(OUT / "photos_5.docx", OUT)

    def test_page_counts(self) -> None:
        self.assertEqual(pdf_page_count(self.pdf_0), 4)
        self.assertEqual(pdf_page_count(self.pdf_3), 5)
        self.assertEqual(pdf_page_count(self.pdf_5), 6)


if __name__ == "__main__":
    unittest.main()
